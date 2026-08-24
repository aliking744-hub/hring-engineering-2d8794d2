from __future__ import annotations

import asyncio
import ipaddress
import re
import shutil
import socket
import subprocess
from collections.abc import Callable
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import cast
from urllib.parse import unquote, urlsplit

import httpx
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text


MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 3_000_000
MAX_OCR_PAGES = 50
ARTICLE_PATTERN = re.compile(r"(?=ماده\s+[\u06F0-\u06F9۰-۹0-9]+)")
ARTICLE_NUMBER_PATTERN = re.compile(r"^ماده\s+([\u06F0-\u06F9۰-۹0-9]+)")
_typed_rtf_to_text = cast(Callable[[str], str], rtf_to_text)


class LegalIngestionError(RuntimeError):
    pass


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    mime_type: str | None
    filename: str | None
    source_type: str
    ocr_used: bool = False


@dataclass(frozen=True)
class LegalTextChunk:
    content: str
    article_number: str | None


def html_to_text(raw_html: str) -> str:
    if len(raw_html.encode("utf-8")) > MAX_DOCUMENT_BYTES:
        raise LegalIngestionError("HTML content exceeds the 10 MB limit")
    soup = BeautifulSoup(raw_html, "html.parser")
    for tag in soup(["script", "style", "head", "nav", "footer", "header", "noscript"]):
        tag.decompose()
    lines = [re.sub(r"\s+", " ", line).strip() for line in soup.get_text("\n").splitlines()]
    text = "\n".join(line for line in lines if line).strip()
    if len(text) < 20:
        raise LegalIngestionError("Not enough textual content was found")
    return text[:MAX_EXTRACTED_CHARACTERS]


def _run_local_command(command: list[str], *, timeout: int = 90) -> str:
    try:
        completed = subprocess.run(
            command,
            check=False,
            capture_output=True,
            timeout=timeout,
            env={"PATH": "/usr/local/sbin:/usr/local/bin:/usr/sbin:/usr/bin:/sbin:/bin", "OMP_THREAD_LIMIT": "1"},
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise LegalIngestionError("Local OCR failed or timed out") from exc
    if completed.returncode != 0:
        raise LegalIngestionError("Local OCR could not read the document")
    return completed.stdout.decode("utf-8", errors="replace")[:MAX_EXTRACTED_CHARACTERS]


def _require_ocr_tools(*tools: str) -> None:
    missing = [tool for tool in tools if shutil.which(tool) is None]
    if missing:
        raise LegalIngestionError("Local OCR runtime is not installed")


def _ocr_image(path: Path) -> str:
    _require_ocr_tools("tesseract")
    return _run_local_command(
        ["tesseract", str(path), "stdout", "-l", "fas+eng", "--psm", "6"]
    ).strip()


def _ocr_pdf(raw: bytes) -> str:
    _require_ocr_tools("pdftoppm", "tesseract")
    with TemporaryDirectory(prefix="hring-legal-ocr-") as directory:
        root = Path(directory)
        input_path = root / "input.pdf"
        input_path.write_bytes(raw)
        prefix = root / "page"
        _run_local_command(
            [
                "pdftoppm",
                "-png",
                "-r",
                "180",
                "-f",
                "1",
                "-l",
                str(MAX_OCR_PAGES),
                str(input_path),
                str(prefix),
            ],
            timeout=120,
        )
        pages = sorted(root.glob("page-*.png"))
        if not pages:
            raise LegalIngestionError("Scanned PDF produced no readable pages")
        text_parts = [_ocr_image(page) for page in pages]
    return "\n\n".join(part for part in text_parts if part).strip()


def _extract_upload_sync(
    filename: str,
    raw: bytes,
    declared_content_type: str | None,
) -> ExtractedDocument:
    if not raw:
        raise LegalIngestionError("Uploaded document is empty")
    if len(raw) > MAX_DOCUMENT_BYTES:
        raise LegalIngestionError("Document exceeds the 10 MB limit")

    safe_name = Path(filename or "document").name
    suffix = Path(safe_name).suffix.lower()
    mime_type = (declared_content_type or "").split(";", 1)[0].strip() or None
    try:
        if suffix == ".pdf" or raw.startswith(b"%PDF-"):
            reader = PdfReader(BytesIO(raw))
            extracted = "\n\n".join(
                (page.extract_text() or "").strip() for page in reader.pages[:MAX_OCR_PAGES]
            ).strip()
            if len(extracted) >= 50:
                return ExtractedDocument(
                    text=extracted[:MAX_EXTRACTED_CHARACTERS],
                    mime_type="application/pdf",
                    filename=safe_name,
                    source_type="upload-pdf",
                )
            ocr_text = _ocr_pdf(raw)
            if len(ocr_text) < 20:
                raise LegalIngestionError("OCR found no usable text in the PDF")
            return ExtractedDocument(
                text=ocr_text,
                mime_type="application/pdf",
                filename=safe_name,
                source_type="upload-pdf-ocr",
                ocr_used=True,
            )
        if suffix == ".docx":
            document = Document(BytesIO(raw))
            text = "\n".join(
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            )
            return ExtractedDocument(
                text=text,
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                filename=safe_name,
                source_type="upload-docx",
            )
        if suffix == ".txt":
            return ExtractedDocument(
                text=raw.decode("utf-8-sig", errors="replace").strip(),
                mime_type="text/plain",
                filename=safe_name,
                source_type="upload-text",
            )
        if suffix == ".rtf":
            return ExtractedDocument(
                text=_typed_rtf_to_text(raw.decode("utf-8", errors="replace")).strip(),
                mime_type="application/rtf",
                filename=safe_name,
                source_type="upload-rtf",
            )
        if suffix == ".doc":
            raise LegalIngestionError(
                "Legacy .doc is not safely parseable; save it as .docx and upload again"
            )
        if suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}:
            with TemporaryDirectory(prefix="hring-legal-image-") as directory:
                image_path = Path(directory) / f"input{suffix}"
                image_path.write_bytes(raw)
                text = _ocr_image(image_path)
            return ExtractedDocument(
                text=text,
                mime_type=mime_type or f"image/{suffix.lstrip('.')}",
                filename=safe_name,
                source_type="upload-image-ocr",
                ocr_used=True,
            )
    except LegalIngestionError:
        raise
    except Exception as exc:
        raise LegalIngestionError("Could not extract text from the uploaded document") from exc
    raise LegalIngestionError("Unsupported document type")


async def extract_upload(
    filename: str,
    raw: bytes,
    declared_content_type: str | None,
) -> ExtractedDocument:
    document = await asyncio.to_thread(
        _extract_upload_sync,
        filename,
        raw,
        declared_content_type,
    )
    if len(document.text.strip()) < 20:
        raise LegalIngestionError("Not enough textual content was found")
    return document


def _forbidden_address(address: str) -> bool:
    ip = ipaddress.ip_address(address)
    return bool(
        ip.is_private
        or ip.is_loopback
        or ip.is_link_local
        or ip.is_multicast
        or ip.is_reserved
        or ip.is_unspecified
    )


async def validate_public_url(url: str) -> str:
    parsed = urlsplit(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise LegalIngestionError("Only public http/https source URLs are allowed")
    if parsed.username or parsed.password:
        raise LegalIngestionError("Source URL credentials are not allowed")
    hostname = parsed.hostname.lower().rstrip(".")
    if hostname in {"localhost", "localhost.localdomain"}:
        raise LegalIngestionError("Local source URLs are not allowed")
    try:
        infos = await asyncio.to_thread(
            socket.getaddrinfo,
            hostname,
            parsed.port or (443 if parsed.scheme == "https" else 80),
        )
    except socket.gaierror as exc:
        raise LegalIngestionError("Source host could not be resolved") from exc
    addresses = {str(item[4][0]) for item in infos}
    if not addresses or any(_forbidden_address(address) for address in addresses):
        raise LegalIngestionError("Private or reserved source addresses are not allowed")
    return parsed.geturl()


async def fetch_public_source(url: str) -> ExtractedDocument:
    safe_url = await validate_public_url(url)
    timeout = httpx.Timeout(20.0, connect=5.0)
    try:
        async with httpx.AsyncClient(timeout=timeout, follow_redirects=False) as client:
            async with client.stream(
                "GET",
                safe_url,
                headers={"User-Agent": "HRing-Legal-Importer/2.0"},
            ) as response:
                if 300 <= response.status_code < 400:
                    raise LegalIngestionError(
                        "Redirects are disabled; provide the final source URL"
                    )
                response.raise_for_status()
                declared = response.headers.get("content-length")
                if declared and int(declared) > MAX_DOCUMENT_BYTES:
                    raise LegalIngestionError("Remote document exceeds the 10 MB limit")
                body = bytearray()
                async for piece in response.aiter_bytes():
                    body.extend(piece)
                    if len(body) > MAX_DOCUMENT_BYTES:
                        raise LegalIngestionError("Remote document exceeds the 10 MB limit")
                content_type = response.headers.get("content-type", "").lower()
    except LegalIngestionError:
        raise
    except (httpx.HTTPError, ValueError) as exc:
        raise LegalIngestionError("Could not retrieve the legal source URL") from exc

    raw = bytes(body)
    if "html" in content_type:
        return ExtractedDocument(
            text=html_to_text(raw.decode("utf-8", errors="replace")),
            mime_type=content_type.split(";", 1)[0],
            filename=None,
            source_type="url-html",
        )
    filename = Path(unquote(urlsplit(safe_url).path)).name or "source.txt"
    extracted = await extract_upload(filename, raw, content_type)
    return ExtractedDocument(
        text=extracted.text,
        mime_type=extracted.mime_type,
        filename=filename,
        source_type=f"url-{extracted.source_type.removeprefix('upload-')}",
        ocr_used=extracted.ocr_used,
    )


def _split_long_section(
    content: str,
    *,
    article_number: str | None,
    chunk_size: int,
    overlap: int,
) -> list[LegalTextChunk]:
    chunks: list[LegalTextChunk] = []
    start = 0
    while start < len(content):
        end = min(start + chunk_size, len(content))
        if end < len(content):
            boundary = content.rfind(" ", start + max(1, chunk_size // 2), end)
            if boundary > start:
                end = boundary
        piece = content[start:end].strip()
        if len(piece) >= 20:
            chunks.append(LegalTextChunk(content=piece, article_number=article_number))
        if end >= len(content):
            break
        next_start = max(start + 1, end - overlap)
        start = next_start
    return chunks


def chunk_legal_text(
    value: str,
    *,
    chunk_size: int,
    overlap: int,
) -> list[LegalTextChunk]:
    normalized = value.replace("\r\n", "\n").replace("\r", "\n").strip()
    if len(normalized) < 20:
        raise LegalIngestionError("Not enough textual content was found")
    if overlap >= chunk_size:
        raise LegalIngestionError("Chunk overlap must be smaller than chunk size")

    sections: list[tuple[str, str | None]] = []
    for part in ARTICLE_PATTERN.split(normalized):
        text = part.strip()
        if len(text) < 20:
            continue
        match = ARTICLE_NUMBER_PATTERN.match(text)
        article_number = match.group(1) if match else ("مقدمه" if not sections else None)
        sections.append((text, article_number))
    if not sections:
        sections = [(normalized, None)]

    chunks: list[LegalTextChunk] = []
    for content, article_number in sections:
        chunks.extend(
            _split_long_section(
                content,
                article_number=article_number,
                chunk_size=chunk_size,
                overlap=overlap,
            )
        )
    if not chunks:
        raise LegalIngestionError("No indexable legal chunks were produced")
    return chunks
